"""Fail-closed generic CSV/XLSX mapping for unsupported bank templates."""

from __future__ import annotations

import csv
import hashlib
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import load_workbook

from app.core.money import MAX_MONEY, money_decimal
from app.core.statement_file_safety import UnsafeStatementFile, validate_xlsx_archive
from app.core.statement_parser import StatementParseResult, StatementTransaction
from app.schemas.statement import MappedImportMapping


MAX_MAPPED_ROWS = 10_000
MAX_MAPPED_COLUMNS = 50
MAX_MAPPED_CELL_LENGTH = 2_000
MAX_MAPPING_ERRORS = 100
GENERIC_MAPPING_VERSION = "1.0"


class MappedImportError(ValueError):
    pass


@dataclass(frozen=True)
class MappedSheet:
    file_format: str
    rows: list[list[Any]]
    header_row: int
    headers: list[str]
    source_fingerprint: str
    header_signature: str


@dataclass(frozen=True)
class MappedRowError:
    row: int
    code: str
    message: str

    def to_dict(self) -> dict[str, object]:
        return {"row": self.row, "code": self.code, "message": self.message}


@dataclass(frozen=True)
class MappedParseResult:
    statement: StatementParseResult
    errors: tuple[MappedRowError, ...]
    identifiers: tuple[str, ...]
    header_signature: str


_HEADER_SYNONYMS = {
    "date_column": ("date", "transaction date", "txn date", "posting date"),
    "value_date_column": ("value date",),
    "description_column": (
        "description",
        "narration",
        "transaction details",
        "remarks",
        "particulars",
    ),
    "reference_column": (
        "reference",
        "reference number",
        "ref no",
        "chq ref no",
        "transaction id",
    ),
    "debit_column": (
        "debit",
        "withdrawal",
        "withdrawal amount",
        "debit amount",
    ),
    "credit_column": (
        "credit",
        "deposit",
        "deposit amount",
        "credit amount",
    ),
    "running_balance_column": (
        "balance",
        "running balance",
        "closing balance",
        "available balance",
    ),
    "account_identifier_column": (
        "account",
        "account number",
        "account no",
        "card number",
    ),
}


def _normalized_header(value: object) -> str:
    text = re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().lower())
    return re.sub(r"\s+", " ", text).strip()


def _bounded_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    if len(text) > MAX_MAPPED_CELL_LENGTH:
        raise MappedImportError(
            f"A spreadsheet cell exceeds {MAX_MAPPED_CELL_LENGTH:,} characters."
        )
    return text


def _read_csv(contents: bytes) -> list[list[Any]]:
    try:
        text = contents.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise MappedImportError(
            "CSV files must use UTF-8 text. Export the sheet as UTF-8 CSV and retry."
        ) from exc
    if "\x00" in text:
        raise MappedImportError("The CSV contains binary data and cannot be mapped.")
    sample = text[:16_384]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows: list[list[Any]] = []
    for row in reader:
        rows.append([_bounded_text(value) for value in row])
        if len(rows) > MAX_MAPPED_ROWS + 20:
            raise MappedImportError(
                f"Mapped imports are limited to {MAX_MAPPED_ROWS:,} transaction rows."
            )
    return rows


def _read_xlsx(contents: bytes) -> list[list[Any]]:
    try:
        validate_xlsx_archive(contents)
    except UnsafeStatementFile as exc:
        raise MappedImportError(str(exc)) from exc
    try:
        workbook = load_workbook(
            io.BytesIO(contents),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise MappedImportError(
            "The XLSX workbook could not be opened safely."
        ) from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise MappedImportError(
                "Choose a workbook containing exactly one transaction sheet."
            )
        rows: list[list[Any]] = []
        for raw_row in workbook.active.iter_rows(values_only=True):
            if len(raw_row) > MAX_MAPPED_COLUMNS:
                raise MappedImportError(
                    f"Mapped imports support at most {MAX_MAPPED_COLUMNS} columns."
                )
            row = []
            for value in raw_row:
                if isinstance(value, str):
                    value = _bounded_text(value)
                row.append(value)
            rows.append(row)
            if len(rows) > MAX_MAPPED_ROWS + 20:
                raise MappedImportError(
                    f"Mapped imports are limited to {MAX_MAPPED_ROWS:,} transaction rows."
                )
        return rows
    finally:
        workbook.close()


def _header_row(rows: list[list[Any]]) -> int:
    candidates: list[tuple[int, int]] = []
    known_headers = {
        synonym
        for synonyms in _HEADER_SYNONYMS.values()
        for synonym in synonyms
    }
    for index, row in enumerate(rows[:20]):
        values = [_normalized_header(value) for value in row]
        populated = [value for value in values if value]
        if len(populated) >= 4 and len(set(populated)) == len(populated):
            matches = sum(value in known_headers for value in populated)
            candidates.append((matches, index))
    if candidates:
        return max(candidates, key=lambda candidate: (candidate[0], -candidate[1]))[1]
    raise MappedImportError(
        "GODFIN could not identify a header row in the first 20 rows."
    )


def inspect_mapped_sheet(contents: bytes, filename: str) -> MappedSheet:
    lower = filename.lower()
    if lower.endswith(".csv"):
        file_format = "csv"
        rows = _read_csv(contents)
    elif lower.endswith(".xlsx"):
        file_format = "xlsx"
        rows = _read_xlsx(contents)
    else:
        raise MappedImportError("Guided mapping supports CSV and XLSX files.")
    if not rows:
        raise MappedImportError("The spreadsheet is empty.")
    header_index = _header_row(rows)
    raw_headers = rows[header_index]
    if len(raw_headers) > MAX_MAPPED_COLUMNS:
        raise MappedImportError(
            f"Mapped imports support at most {MAX_MAPPED_COLUMNS} columns."
        )
    headers = [
        _bounded_text(value) or f"Unnamed column {index + 1}"
        for index, value in enumerate(raw_headers)
    ]
    normalized = [_normalized_header(value) for value in headers]
    signature = hashlib.sha256(
        "\x1f".join(normalized).encode("utf-8")
    ).hexdigest()
    return MappedSheet(
        file_format=file_format,
        rows=rows,
        header_row=header_index,
        headers=headers,
        source_fingerprint=hashlib.sha256(contents).hexdigest(),
        header_signature=signature,
    )


def suggested_mapping(headers: list[str]) -> dict[str, int | None]:
    normalized = [_normalized_header(header) for header in headers]
    result: dict[str, int | None] = {}
    for field, synonyms in _HEADER_SYNONYMS.items():
        matches = [
            index
            for index, header in enumerate(normalized)
            if header in synonyms
        ]
        result[field] = matches[0] if len(matches) == 1 else None
    return result


def mapping_fingerprint(
    sheet: MappedSheet,
    mapping: MappedImportMapping,
    date_format: str,
) -> str:
    canonical = mapping.model_dump_json(exclude_none=False)
    return hashlib.sha256(
        f"{sheet.source_fingerprint}\x1f{date_format}\x1f{canonical}".encode("utf-8")
    ).hexdigest()


def _mapping_columns(mapping: MappedImportMapping) -> list[int]:
    return [
        value
        for field, value in mapping.model_dump().items()
        if field.endswith("_column") and value is not None
    ]


def _validate_mapping(mapping: MappedImportMapping, column_count: int) -> None:
    columns = _mapping_columns(mapping)
    if any(column >= column_count for column in columns):
        raise MappedImportError("A selected mapping column does not exist.")
    required = {
        mapping.date_column,
        mapping.description_column,
        mapping.debit_column,
        mapping.credit_column,
    }
    if len(required) != 4:
        raise MappedImportError(
            "Date, description, debit, and credit must use four different columns."
        )
    if len(columns) != len(set(columns)):
        raise MappedImportError("Each spreadsheet column can be mapped only once.")


def _cell(row: list[Any], column: int | None) -> Any:
    if column is None or column >= len(row):
        return None
    return row[column]


def _parse_date(value: object, date_format: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = _bounded_text(value)
    formats = {
        "dd/mm/yyyy": "%d/%m/%Y",
        "dd-mm-yyyy": "%d-%m-%Y",
        "yyyy-mm-dd": "%Y-%m-%d",
        "mm/dd/yyyy": "%m/%d/%Y",
    }
    if date_format != "auto":
        try:
            return datetime.strptime(text, formats[date_format]).date()
        except (KeyError, ValueError) as exc:
            raise MappedImportError(
                f"Date does not match the selected {date_format} format."
            ) from exc
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
        return datetime.strptime(text, "%Y-%m-%d").date()
    match = re.fullmatch(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", text)
    if not match:
        raise MappedImportError("Date format is not recognized.")
    first, second = int(match.group(1)), int(match.group(2))
    if first <= 12 and second <= 12:
        raise MappedImportError(
            "Date is ambiguous. Select day/month/year or month/day/year explicitly."
        )
    fmt = "%d/%m/%Y" if "/" in text and first > 12 else None
    if "-" in text and first > 12:
        fmt = "%d-%m-%Y"
    if "/" in text and second > 12:
        fmt = "%m/%d/%Y"
    if fmt is None:
        raise MappedImportError("Date is outside the supported calendar format.")
    try:
        return datetime.strptime(text, fmt).date()
    except ValueError as exc:
        raise MappedImportError("Date is not a valid calendar day.") from exc


def _parse_money(value: object, *, allow_negative: bool = False) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, bool):
        raise MappedImportError("Amount is not numeric.")
    if isinstance(value, (int, float, Decimal)):
        raw = str(value)
    else:
        raw = _bounded_text(value)
    text = raw.strip()
    if not text or text in {"-", "--"}:
        return None
    negative_parentheses = text.startswith("(") and text.endswith(")")
    text = text.strip("()").replace(",", "")
    text = re.sub(r"^(?:₹|INR|Rs\.?)\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\s*(?:Cr|Dr)$", "", text, flags=re.IGNORECASE)
    try:
        amount = Decimal(text)
    except (InvalidOperation, ValueError) as exc:
        raise MappedImportError("Amount is not numeric.") from exc
    if negative_parentheses:
        amount = -amount
    if not amount.is_finite() or abs(amount) > MAX_MONEY:
        raise MappedImportError("Amount is outside GODFIN's supported range.")
    amount = money_decimal(amount)
    if amount < 0 and not allow_negative:
        raise MappedImportError("Debit and credit columns must use positive amounts.")
    return amount


def parse_mapped_sheet(
    sheet: MappedSheet,
    mapping: MappedImportMapping,
    *,
    date_format: str,
) -> MappedParseResult:
    _validate_mapping(mapping, len(sheet.headers))
    transactions: list[StatementTransaction] = []
    errors: list[MappedRowError] = []
    identifiers: set[str] = set()

    for row_index, row in enumerate(
        sheet.rows[sheet.header_row + 1 :],
        start=sheet.header_row + 2,
    ):
        if all(not _bounded_text(value) for value in row if value is not None):
            continue
        try:
            transaction_date = _parse_date(
                _cell(row, mapping.date_column), date_format
            )
            description = _bounded_text(
                _cell(row, mapping.description_column)
            )
            if not description:
                raise MappedImportError("Description is empty.")
            debit = _parse_money(_cell(row, mapping.debit_column))
            credit = _parse_money(_cell(row, mapping.credit_column))
            has_debit = debit is not None and debit > 0
            has_credit = credit is not None and credit > 0
            if has_debit == has_credit:
                raise MappedImportError(
                    "Exactly one of debit or credit must contain a positive amount."
                )
            transaction_type = "debit" if has_debit else "credit"
            amount = debit if has_debit else credit
            assert amount is not None
            reference = _bounded_text(
                _cell(row, mapping.reference_column)
            ) or None
            value_date = None
            if mapping.value_date_column is not None:
                raw_value_date = _cell(row, mapping.value_date_column)
                if _bounded_text(raw_value_date):
                    value_date = _parse_date(raw_value_date, date_format)
            balance = None
            if mapping.running_balance_column is not None:
                balance = _parse_money(
                    _cell(row, mapping.running_balance_column),
                    allow_negative=True,
                )
                if balance is None:
                    raise MappedImportError("Running balance is empty.")
            identifier = _bounded_text(
                _cell(row, mapping.account_identifier_column)
            )
            if identifier:
                identifiers.add(identifier)
            transactions.append(
                StatementTransaction(
                    date=transaction_date,
                    value_date=value_date,
                    description=description,
                    amount=float(amount),
                    txn_type=transaction_type,
                    ref_number=reference,
                    closing_balance=float(balance) if balance is not None else None,
                    instrument="mapped_spreadsheet",
                    semantic_type="expense" if transaction_type == "debit" else "unknown",
                )
            )
        except MappedImportError as exc:
            if len(errors) < MAX_MAPPING_ERRORS:
                errors.append(
                    MappedRowError(
                        row=row_index,
                        code="INVALID_MAPPED_ROW",
                        message=str(exc),
                    )
                )

    if len(identifiers) > 1:
        errors.append(
            MappedRowError(
                row=sheet.header_row + 1,
                code="MULTIPLE_ACCOUNT_IDENTIFIERS",
                message="The file contains more than one account identifier.",
            )
        )
    if not transactions and not errors:
        errors.append(
            MappedRowError(
                row=sheet.header_row + 1,
                code="NO_TRANSACTION_ROWS",
                message="No transaction rows were found below the header.",
            )
        )
    if len(transactions) > MAX_MAPPED_ROWS:
        raise MappedImportError(
            f"Mapped imports are limited to {MAX_MAPPED_ROWS:,} transactions."
        )

    opening_balance = None
    closing_balance = None
    reconciliation_status = "mapped_valid"
    reconciliation_method = "explicit_user_column_mapping"
    if mapping.running_balance_column is not None and transactions and not errors:
        ascending = all(
            left.date <= right.date
            for left, right in zip(transactions, transactions[1:])
        )
        descending = all(
            left.date >= right.date
            for left, right in zip(transactions, transactions[1:])
        )
        if not ascending and not descending:
            errors.append(
                MappedRowError(
                    row=sheet.header_row + 2,
                    code="UNORDERED_RUNNING_BALANCE",
                    message=(
                        "Rows with a running balance must be ordered by date "
                        "from oldest to newest or newest to oldest."
                    ),
                )
            )
        else:
            ordered = transactions if ascending else list(reversed(transactions))
            first_balance = money_decimal(ordered[0].closing_balance)
            first_amount = money_decimal(ordered[0].amount)
            first_signed = (
                first_amount if ordered[0].txn_type == "credit" else -first_amount
            )
            opening_balance = (first_balance - first_signed).quantize(
                Decimal("0.01")
            )
            previous = opening_balance
            for transaction in ordered:
                amount = money_decimal(transaction.amount)
                signed = amount if transaction.txn_type == "credit" else -amount
                expected = (previous + signed).quantize(Decimal("0.01"))
                actual = money_decimal(transaction.closing_balance)
                if expected != actual:
                    errors.append(
                        MappedRowError(
                            row=sheet.header_row + 2,
                            code="RUNNING_BALANCE_MISMATCH",
                            message=(
                                "Running-balance arithmetic does not match the "
                                "mapped debit and credit columns."
                            ),
                        )
                    )
                    break
                previous = actual
            if not errors:
                closing_balance = previous
                reconciliation_status = "passed"
                reconciliation_method = "mapped_columns_and_running_balance"

    statement = StatementParseResult(
        transactions=transactions,
        statement_type="generic_mapped",
        parser_profile="generic_mapped",
        recognized=True,
        reconciliation_status=("failed" if errors else reconciliation_status),
        reconciliation_method=reconciliation_method,
        source_digest=sheet.source_fingerprint,
        period_start=min((item.date for item in transactions), default=None),
        period_end=max((item.date for item in transactions), default=None),
        opening_balance=(float(opening_balance) if opening_balance is not None else None),
        closing_balance=(float(closing_balance) if closing_balance is not None else None),
        total_debits=float(
            sum(
                (money_decimal(item.amount) for item in transactions if item.txn_type == "debit"),
                Decimal("0.00"),
            )
        ),
        total_credits=float(
            sum(
                (money_decimal(item.amount) for item in transactions if item.txn_type == "credit"),
                Decimal("0.00"),
            )
        ),
        errors=[error.message for error in errors],
    )
    return MappedParseResult(
        statement=statement,
        errors=tuple(errors),
        identifiers=tuple(sorted(identifiers)),
        header_signature=sheet.header_signature,
    )


def sample_rows(sheet: MappedSheet, limit: int = 5) -> list[list[str]]:
    rows = []
    for raw in sheet.rows[sheet.header_row + 1 :]:
        if not any(_bounded_text(value) for value in raw if value is not None):
            continue
        rows.append([_bounded_text(value) for value in raw[: len(sheet.headers)]])
        if len(rows) >= limit:
            break
    return rows
